"""Macro workbook resources close on every edit/preview exit path."""

from unittest.mock import Mock
from zipfile import ZipFile

import openpyxl
import pytest

from app.core import db
from app.services import document_intelligence_service as service


@pytest.mark.parametrize("mode", ["preview", "write", "missing_sheet", "save_error"])
def test_macro_edit_closes_vba_archive_without_losing_content(monkeypatch, tmp_path, mode):
    db.init_db()
    source, target = tmp_path / "source.xlsx", tmp_path / "macro.xlsm"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "original"
    workbook.save(source)
    workbook.close()
    with ZipFile(source) as original, ZipFile(target, "w") as archive:
        for name in original.namelist():
            archive.writestr(name, original.read(name))
        archive.writestr("xl/vbaProject.bin", b"SYNTHETIC_VBA_PART")
    before = target.read_bytes()
    loaded = []
    real_load = openpyxl.load_workbook

    def capture_load(*args, **kwargs):
        result = real_load(*args, **kwargs)
        loaded.append(result)
        monkeypatch.setattr(result, "close", Mock(wraps=result.close))
        if mode == "save_error":

            def fail_save(*args, **kwargs):
                raise OSError("injected save failure")

            monkeypatch.setattr(result, "save", fail_save)
        return result

    monkeypatch.setattr(openpyxl, "load_workbook", capture_load)
    kwargs = dict(
        sheet="missing" if mode == "missing_sheet" else "Sheet",
        cell="A1",
        value="updated",
        dry_run=mode in {"preview", "missing_sheet"},
    )
    if mode in {"missing_sheet", "save_error"}:
        with pytest.raises(ValueError if mode == "missing_sheet" else OSError):
            service.edit_xlsx(target, **kwargs)
    else:
        service.edit_xlsx(target, **kwargs)

    assert len(loaded) == 1
    assert loaded[0].vba_archive.fp is None
    loaded[0].close.assert_called_once_with()
    if mode != "write":
        assert target.read_bytes() == before
    with ZipFile(target) as archive:
        assert archive.read("xl/vbaProject.bin") == b"SYNTHETIC_VBA_PART"
    result = real_load(target, read_only=True)
    try:
        assert result.active["A1"].value == ("updated" if mode == "write" else "original")
    finally:
        result.close()
