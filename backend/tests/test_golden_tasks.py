"""黄金任务回归套件（golden tasks）。

数据集：``test_data/golden_tasks/golden_tasks.json``。
目标：用真实 API / 工具路径断言关键产物（计划、风险、审批、文件副作用、
工具输出），而不是仅断言返回码；全部任务离线确定性可跑（MockProvider /
确定性规划器 / extractive fallback）。

证据边界：本套件是机器自证的版本回归证据，不能替代真人对自然语言结果
质量（成功率 / 可读性 / 返工率）的评分签收；真人评审流程见
``docs/qa/golden-tasks.md`` 与 ``npm run evidence:result-quality-review``。
"""

from __future__ import annotations

import json
import os
import time
from pathlib import Path
from typing import Any

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.routes_approvals import router as approvals_router
from app.api.routes_chat import router as chat_router
from app.api.routes_files import router as files_router
from app.api.routes_runs import router as runs_router
from app.core import db
from app.core.errors import SecurityError
from app.services import run_service
from app.tools.registry import register_all_tools, registry

PROJECT_ROOT = Path(__file__).resolve().parents[2]
GOLDEN_DATASET_PATH = PROJECT_ROOT / "test_data" / "golden_tasks" / "golden_tasks.json"
GOLDEN_DATASET = json.loads(GOLDEN_DATASET_PATH.read_text(encoding="utf-8"))
GOLDEN_TASKS: list[dict[str, Any]] = GOLDEN_DATASET["tasks"]

TERMINAL_OR_WAITING = {"completed", "failed", "denied", "cancelled", "awaiting_approval"}
ERROR_CLASSES = {"SecurityError": SecurityError}


def test_golden_dataset_integrity():
    """数据集本身是发布门禁的一部分：数量、唯一性、必填字段。"""
    ids = [task["id"] for task in GOLDEN_TASKS]
    assert len(ids) == len(set(ids)), "golden task id must be unique"
    assert len(ids) >= 30, f"golden task count must stay >= 30, got {len(ids)}"
    categories = {task["category"] for task in GOLDEN_TASKS}
    for required in ("system", "cleanup", "approval", "safety", "file", "document", "chat", "browser", "search"):
        assert required in categories, f"missing golden category: {required}"
    for task in GOLDEN_TASKS:
        assert task.get("entry") in {"runs", "chat", "files_api", "tool"}, task["id"]
        assert task.get("expect"), task["id"]


def test_golden_path_placeholders_use_native_separators(tmp_path: Path):
    workspace = tmp_path / "workspace"
    outside = tmp_path / "outside-secret.txt"

    assert _sub("$WS\\old-report.txt", workspace, outside) == str(workspace / "old-report.txt")
    assert _sub("$WS/old-report.txt", workspace, outside) == str(workspace / "old-report.txt")
    assert _sub({"path": "$OUTSIDE"}, workspace, outside) == {"path": str(outside)}


@pytest.mark.parametrize("task", GOLDEN_TASKS, ids=[t["id"] for t in GOLDEN_TASKS])
def test_golden_task(task, monkeypatch, tmp_path):
    workspace, outside = _golden_env(monkeypatch, tmp_path, task)
    handler = {
        "runs": _run_entry,
        "chat": _chat_entry,
        "files_api": _files_api_entry,
        "tool": _tool_entry,
    }[task["entry"]]
    handler(task, workspace, outside)


# ---------------------------------------------------------------------------
# environment


def _write_docx_fixture(target: Path, spec: dict[str, Any]) -> None:
    from docx import Document

    doc = Document()
    for paragraph in spec.get("paragraphs") or []:
        doc.add_paragraph(str(paragraph))
    doc.save(str(target))


def _golden_env(monkeypatch, tmp_path, task: dict[str, Any]):
    workspace = tmp_path / "workspace"
    workspace.mkdir(parents=True, exist_ok=True)
    outside = tmp_path / "outside-secret.txt"
    outside.write_text("outside the authorized scope", encoding="utf-8")
    for rel, content in (task.get("fixtures") or {}).items():
        target = workspace / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(content, dict) and content.get("__docx__"):
            _write_docx_fixture(target, content)
        else:
            target.write_text(content, encoding="utf-8")

    monkeypatch.setenv("LENGRVIS_CONFIG_FILE", str(tmp_path / "missing-config.yaml"))
    monkeypatch.setenv("LENGRVIS_ENV_FILE", str(tmp_path / "missing.env"))
    monkeypatch.setenv("LENGRVIS_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("LENGRVIS_PROVIDER_NAME", "mock")
    monkeypatch.setenv("LENGRVIS_API_KEY", "")
    monkeypatch.setenv("LENGRVIS_ALLOW_MOCK_FALLBACK", "true")
    monkeypatch.setenv("LENGRVIS_MODE", task.get("mode", "efficiency"))
    if task.get("no_scope"):
        monkeypatch.delenv("LENGRVIS_ALLOWED_DIRECTORIES", raising=False)
    else:
        monkeypatch.setenv("LENGRVIS_ALLOWED_DIRECTORIES", str(workspace))
    monkeypatch.setattr("app.llm.registry.detect_local_backend", lambda: None)
    monkeypatch.setattr("app.llm.registry.detect_onnx_backend", lambda settings=None: None)
    db.init_db()
    return workspace, outside


def _sub(value: Any, workspace: Path, outside: Path) -> Any:
    if isinstance(value, str):
        return _replace_path_placeholder(_replace_path_placeholder(value, "$WS", workspace), "$OUTSIDE", outside)
    if isinstance(value, dict):
        return {key: _sub(item, workspace, outside) for key, item in value.items()}
    if isinstance(value, list):
        return [_sub(item, workspace, outside) for item in value]
    return value


def _replace_path_placeholder(value: str, placeholder: str, path: Path) -> str:
    text = value
    for separator in ("\\", "/"):
        text = text.replace(f"{placeholder}{separator}", str(path) + os.sep)
    return text.replace(placeholder, str(path))


def _register_golden_deferred_tools(task: dict[str, Any]) -> None:
    """tool.search only indexes deferred tools; golden tasks may register stubs."""
    from app.policy.risk import RiskLevel
    from app.tools.schemas import ToolDefinition

    for spec in task.get("deferred_tools") or []:
        registry.register(
            ToolDefinition(
                name=str(spec["name"]),
                description=str(spec.get("description") or spec["name"].replace(".", " ")),
                input_schema={"type": "object", "properties": {}},
                output_schema={"type": "object", "properties": {"ok": {"type": "boolean"}}},
                risk_level=RiskLevel.R0_READ_ONLY,
                agent_owner=str(spec.get("agent_owner") or "SearchAgent"),
                supports_dry_run=False,
                requires_authorized_path=False,
                execute=lambda args, context: {"ok": True},
                search_hint=str(spec.get("search_hint") or spec["name"]),
                defer_loading=True,
                read_only=True,
            )
        )


def _golden_app() -> FastAPI:
    """轻量测试应用：只挂载黄金任务用到的公共路由，避免 lifespan watcher。"""
    app = FastAPI()
    app.include_router(runs_router, prefix="/api")
    app.include_router(approvals_router, prefix="/api")
    app.include_router(chat_router, prefix="/api")
    app.include_router(files_router, prefix="/api")
    return app


# ---------------------------------------------------------------------------
# entry handlers


def _run_entry(task: dict[str, Any], workspace: Path, outside: Path) -> None:
    expect = task["expect"]
    message = _sub(task["message"], workspace, outside)
    with TestClient(_golden_app()) as client:
        created = client.post(
            "/api/runs",
            json={"message": message, "mode": task.get("mode", "efficiency"), "engine": task.get("engine", "os")},
        )
        assert created.status_code == 200, created.text
        run = created.json()
        if task.get("routing_only"):
            _assert_routing_expectations(run, expect)
            return
        final = _wait_for_phase(client, run["run_id"], set(expect["phase"]))
        assert final["phase"] in expect["phase"], f"phase={final['phase']} expected={expect['phase']}"
        task_id = final["task_id"]
        _assert_run_expectations(client, run["run_id"], task_id, expect, workspace)

        after = task.get("after")
        if after:
            approval = _latest_pending_approval(task_id)
            assert approval is not None, "expected a pending approval before the follow-up action"
            response = client.post(f"/api/approvals/{approval['id']}/{after}")
            assert response.status_code == 200, response.text
            after_expect = task["after_expect"]
            final_after = _wait_for_phase(client, run["run_id"], set(after_expect["phase"]))
            assert final_after["phase"] in after_expect["phase"], (
                f"after-phase={final_after['phase']} expected={after_expect['phase']}"
            )
            if "approval_status" in after_expect:
                latest = db.fetch_one("approvals", approval["id"])
                assert latest["status"] == after_expect["approval_status"], latest["status"]
            _assert_file_states(after_expect, workspace)


def _chat_entry(task: dict[str, Any], workspace: Path, outside: Path) -> None:
    expect = task["expect"]
    message = _sub(task["message"], workspace, outside)
    with TestClient(_golden_app()) as client:
        response = client.post("/api/chat", json={"message": message, "mode": task.get("mode", "efficiency")})
        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload["delegated"] == expect["delegated"], payload
        if expect.get("reply_contains"):
            assert expect["reply_contains"] in str(payload.get("message") or ""), payload.get("message")
        if expect.get("no_tasks"):
            assert db.fetch_many("tasks") == []
        if expect.get("agent"):
            assert payload.get("agent") == expect["agent"], payload
        if expect.get("task_plan_tools"):
            task_id = payload.get("task_id")
            assert task_id, payload
            _wait_for_plan_tools(task_id, expect["task_plan_tools"])
        if expect.get("task_metadata_hint"):
            task_id = payload.get("task_id")
            assert task_id, payload
            _wait_for_task_metadata_hint(task_id, expect["task_metadata_hint"])
        if expect.get("task_completed"):
            task_id = payload["task_id"]
            assert task_id
            final = _wait_for_task_status(task_id, "completed", "failed", "denied")
            assert final["status"] == "completed", final["status"]
            if expect.get("task_plan_tools"):
                assert _plan_tools(task_id) == expect["task_plan_tools"]


def _files_api_entry(task: dict[str, Any], workspace: Path, outside: Path) -> None:
    expect = task["expect"]
    with TestClient(_golden_app()) as client:
        response = client.get("/api/files/search", params={"q": task["query"]})
        assert response.status_code == 200, response.text
        payload = response.json()
        name_search = payload.get("name_search") or {}
        if "status" in expect:
            assert name_search.get("status") == expect["status"], name_search
            return
        if "name_count" in expect:
            assert name_search.get("count") == expect["name_count"], name_search
        if expect.get("name_contains"):
            names = {item.get("name") for item in payload.get("name_results") or []}
            assert expect["name_contains"] in names, names


def _tool_entry(task: dict[str, Any], workspace: Path, outside: Path) -> None:
    expect = task["expect"]
    register_all_tools(load_skills=False)
    _register_golden_deferred_tools(task)
    definition = registry.get(task["tool"])
    args = _sub(task.get("args") or {}, workspace, outside)
    settings = __import__("app.llm.registry", fromlist=["get_effective_settings"]).get_effective_settings()
    context = {
        "allowed_directories": [] if task.get("no_scope") else [str(workspace)],
        "settings": settings,
    }

    if expect.get("error"):
        with pytest.raises(ERROR_CLASSES[expect["error"]]):
            definition.execute(args, context)
        return

    output = definition.execute(args, context)
    for spec in expect.get("output") or []:
        _assert_output_spec(output, spec)
    _assert_file_states(expect, workspace)


# ---------------------------------------------------------------------------
# assertions


def _assert_routing_expectations(run: dict[str, Any], expect: dict[str, Any]) -> None:
    if expect.get("route_rule"):
        assert run.get("engine_route_rule") == expect["route_rule"], run
    if expect.get("engine"):
        assert run.get("engine") == expect["engine"], run
    caps = run.get("engine_capabilities") or {}
    for key, value in (expect.get("engine_capabilities") or {}).items():
        assert caps.get(key) == value, f"engine_capabilities[{key}]={caps.get(key)!r} expected {value!r}"


def _assert_run_expectations(
    client: TestClient,
    run_id: str,
    task_id: str,
    expect: dict[str, Any],
    workspace: Path,
) -> None:
    if expect.get("route_rule") or expect.get("engine_capabilities"):
        detail = client.get(f"/api/runs/{run_id}")
        assert detail.status_code == 200, detail.text
        _assert_routing_expectations(detail.json(), expect)
    if expect.get("plan_tools"):
        assert _plan_tools(task_id) == expect["plan_tools"]
    if expect.get("global_risk"):
        plans = db.fetch_many("plans", "task_id = ?", (task_id,), limit=1)
        assert plans, "expected a stored plan"
        assert plans[0]["global_risk_level"] == expect["global_risk"], plans[0]["global_risk_level"]
    if "pending_approvals" in expect:
        pending = [
            approval
            for approval in db.fetch_many("approvals", "task_id = ?", (task_id,), limit=20)
            if approval["status"] == "pending"
        ]
        assert len(pending) == expect["pending_approvals"], f"pending approvals: {len(pending)}"
    if expect.get("no_tool_results"):
        calls = db.fetch_many("tool_calls", "task_id = ?", (task_id,), limit=20)
        results = []
        for call in calls:
            results.extend(db.fetch_many("tool_results", "tool_call_id = ?", (call["id"],), limit=5))
        assert results == [], f"forbidden task must not execute tools: {[c['tool_name'] for c in calls]}"
    for spec in expect.get("tool_output_has_keys") or []:
        outputs = _tool_outputs(task_id, spec["tool"])
        assert outputs, f"no tool output recorded for {spec['tool']}"
        for key in spec["keys"]:
            assert key in outputs[0], f"missing key {key} in {spec['tool']} output: {sorted(outputs[0])}"
    for spec in expect.get("tool_output_contains") or []:
        outputs = _tool_outputs(task_id, spec["tool"])
        assert outputs, f"no tool output recorded for {spec['tool']}"
        encoded = json.dumps(outputs, ensure_ascii=False)
        assert spec["text"] in encoded, f"{spec['text']!r} not in {spec['tool']} output"
    for name in expect.get("timeline_any") or []:
        # The run phase (derived from the task row) can be observed slightly
        # before the resident engine loop publishes the matching run event;
        # poll briefly instead of asserting a single snapshot.
        names: list[str] = []
        for _ in range(40):
            timeline = client.get(f"/api/runs/{run_id}/timeline").json()
            names = [event["name"] for event in timeline.get("events", [])]
            if name in names:
                break
            time.sleep(0.05)
        assert name in names, f"{name} not in timeline events: {names}"
    _assert_file_states(expect, workspace)


def _assert_file_states(expect: dict[str, Any], workspace: Path) -> None:
    for rel in expect.get("files_exist") or []:
        assert (workspace / rel).exists(), f"expected file to survive: {rel}"
    for rel in expect.get("files_absent") or []:
        assert not (workspace / rel).exists(), f"expected file to be gone: {rel}"


def _assert_output_spec(output: Any, spec: dict[str, Any]) -> None:
    value = _get_path(output, spec["path"])
    if "equals" in spec:
        assert value == spec["equals"], f"{spec['path']}={value!r} expected {spec['equals']!r}"
    if "contains" in spec:
        assert spec["contains"] in str(value), f"{spec['contains']!r} not in {spec['path']}={value!r}"
    if "min" in spec:
        assert value is not None and float(value) >= spec["min"], f"{spec['path']}={value!r} < {spec['min']}"
    if "min_len" in spec:
        assert value is not None and len(str(value)) >= spec["min_len"], f"{spec['path']} too short: {value!r}"
    if spec.get("nonempty"):
        assert value, f"{spec['path']} must be non-empty"
    if "json_contains" in spec:
        encoded = json.dumps(value, ensure_ascii=False, default=str)
        assert spec["json_contains"] in encoded, f"{spec['json_contains']!r} not in {spec['path']}"


def _get_path(payload: Any, dotted: str) -> Any:
    value = payload
    for part in dotted.split("."):
        if isinstance(value, dict):
            value = value.get(part)
        else:
            return None
    return value


# ---------------------------------------------------------------------------
# db helpers


def _plan_tools(task_id: str) -> list[str]:
    plans = db.fetch_many("plans", "task_id = ?", (task_id,), limit=1)
    assert plans, "expected a stored plan"
    return [step["tool_name"] for step in plans[0]["steps"]]


def _tool_outputs(task_id: str, tool_name: str) -> list[dict[str, Any]]:
    outputs: list[dict[str, Any]] = []
    for call in db.fetch_many("tool_calls", "task_id = ?", (task_id,), limit=50):
        if call["tool_name"] != tool_name:
            continue
        for result in db.fetch_many("tool_results", "tool_call_id = ?", (call["id"],), limit=5):
            output = result.get("output")
            if isinstance(output, dict):
                outputs.append(output)
    return outputs


def _latest_pending_approval(task_id: str) -> dict[str, Any] | None:
    pending = [a for a in db.fetch_many("approvals", "task_id = ?", (task_id,), limit=20) if a["status"] == "pending"]
    return pending[0] if pending else None


def _wait_for_phase(client: TestClient, run_id: str, target_phases: set[str]) -> dict[str, Any]:
    stop_phases = target_phases | TERMINAL_OR_WAITING
    payload: dict[str, Any] = {}
    for _ in range(240):
        response = client.get(f"/api/runs/{run_id}")
        assert response.status_code == 200
        payload = response.json()
        if payload["phase"] in stop_phases:
            if payload["phase"] in {"completed", "failed", "denied", "cancelled"}:
                _wait_for_run_inactive(run_id)
            return payload
        time.sleep(0.05)
    raise AssertionError(f"Run {run_id} stuck in phase {payload.get('phase')}, expected {target_phases}")


def _wait_for_run_inactive(run_id: str) -> None:
    for _ in range(240):
        if run_id not in run_service.active_run_ids():
            return
        time.sleep(0.05)
    raise AssertionError(f"Run {run_id} still active after terminal phase")


def _wait_for_plan_tools(task_id: str, expected: list[str]) -> None:
    last: list[str] = []
    for _ in range(240):
        plans = db.fetch_many("plans", "task_id = ?", (task_id,), limit=1)
        if plans:
            last = [step["tool_name"] for step in plans[0]["steps"]]
            if last == expected:
                return
        time.sleep(0.05)
    raise AssertionError(f"task {task_id} plan tools {last!r} != expected {expected!r}")


def _wait_for_task_metadata_hint(task_id: str, expected_hint: str) -> None:
    for _ in range(240):
        task = db.fetch_one("tasks", task_id)
        metadata = (task or {}).get("metadata") or {}
        if isinstance(metadata, str):
            metadata = json.loads(metadata) if metadata.startswith("{") else {}
        if metadata.get("supervisor_agent_hint") == expected_hint:
            return
        time.sleep(0.05)
    raise AssertionError(f"task {task_id} metadata hint != {expected_hint!r}")


def _wait_for_task_status(task_id: str, *statuses: str) -> dict[str, Any]:
    task: dict[str, Any] | None = None
    for _ in range(240):
        task = db.fetch_one("tasks", task_id)
        if task and task["status"] in statuses:
            return task
        time.sleep(0.05)
    raise AssertionError(f"Task {task_id} did not reach {statuses}: {task and task.get('status')}")
