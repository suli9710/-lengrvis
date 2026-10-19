"""Regression coverage for the October review's privacy and lifecycle boundaries."""

import json
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest

from app.config import AppSettings
from app.context.management import project_messages_for_llm
from app.core import butler_store as store
from app.core import db
from app.core.butler_models import AutonomyScope, ButlerMatter, MatterInteraction, MatterStatus
from app.core.schemas import Approval, Task, TaskPhase
from app.core.session_context import SessionContextStore
from app.orchestration.butler_runtime import ButlerExecutionDeferred, validate_task_action
from app.policy.permissions import PermissionStore
from app.policy.risk import RiskLevel
from app.services.butler_state import public_matter
from app.services.local_retention_service import cleanup_expired_task_details


@pytest.fixture(autouse=True)
def _database():
    db.init_db()


def test_private_summary_cannot_enter_cloud_projection_after_reload():
    from app.orchestration.handlers.completion_handler import CompletionHandler

    session = SessionContextStore()
    session.load()
    task = Task(user_goal="PRIVATE_SUMMARY_MARKER", mode="privacy")
    bus = SimpleNamespace(
        get_messages=lambda task_id: [SimpleNamespace(id="message", created_at=datetime.now(UTC).isoformat())],
        get_llm_messages=lambda *args, **kwargs: [{"role": "user", "content": task.user_goal}],
    )
    CompletionHandler(SimpleNamespace(bus=bus, session_context_store=session))._update_session_summary(task)
    reloaded = SessionContextStore()
    reloaded.load()
    projection = project_messages_for_llm(
        [{"role": "user", "content": "public follow-up"}],
        AppSettings(mode="efficiency"),
        session_context=reloaded.planning_context(),
    )
    assert "PRIVATE_SUMMARY_MARKER" not in str(projection.messages)
    assert "PRIVATE_SUMMARY_MARKER" in reloaded.current.conversation_summary


def test_summary_source_policy_is_authenticated_and_stays_local_after_private_merge():
    from app.core.context_egress import filter_session_context

    session = SessionContextStore()
    session.load()
    session.remember_summary("public", source_mode="efficiency")
    assert filter_session_context(session.planning_context(), "hybrid")["conversation_summary"] == "public"
    copied = {**session.planning_context(), "session_id": "other"}
    assert not filter_session_context(copied, "efficiency")["conversation_summary"]
    session.remember_summary("public and private", source_mode="privacy")
    session.remember_summary("public and private and later", source_mode="efficiency")
    assert not filter_session_context(session.planning_context(), "efficiency")["conversation_summary"]
    assert filter_session_context(session.planning_context(), "privacy")["conversation_summary"]


@pytest.mark.parametrize("consent", [False, True])
def test_image_provider_boundaries_require_upload_consent(monkeypatch, tmp_path, consent):
    from PIL import Image

    from app.indexer import ocr_service
    from app.tools import vision_tools

    image = tmp_path / "scan.png"
    Image.new("RGB", (2, 2), "white").save(image)
    provider = SimpleNamespace(
        name="cloud",
        vision=AsyncMock(return_value="description"),
        ocr=AsyncMock(return_value="text"),
        embed_image=AsyncMock(return_value=[1.0, 0.0]),
    )
    monkeypatch.setattr(vision_tools, "get_provider", lambda *args, **kwargs: provider)
    monkeypatch.setattr(ocr_service, "get_provider", lambda *args, **kwargs: provider)
    monkeypatch.setattr(ocr_service, "local_ocr_image", lambda *args, **kwargs: ocr_service.OCRResult(ok=False))
    settings = AppSettings(mode="efficiency", allow_file_content_upload=consent)
    vision_tools.describe_image({"path": str(image)}, {"settings": settings, "allowed_directories": [str(tmp_path)]})
    ocr_service.ocr_image_result(image, settings=settings, allow_cloud_fallback=True)
    assert provider.vision.call_count == int(consent)
    assert provider.ocr.call_count == int(consent)
    monkeypatch.setattr(vision_tools, "_local_image_embedding", lambda *args, **kwargs: {"ok": False})
    vision_tools.image_embedding(
        image,
        profile={"description": "private labels"},
        context={"settings": settings, "image_embedding_provider": provider},
    )
    assert provider.embed_image.call_count == int(consent)
    comparison = vision_tools.compare_images(
        {"path_a": str(image), "path_b": str(image)}, {"settings": settings, "allowed_directories": [str(tmp_path)]}
    )
    assert comparison["ok"] is consent


def test_erase_settings_restores_safe_policy_without_weakening_missing_record_guard():
    db.bootstrap_sensitive_record_integrity()
    db.erase_local_user_data(include_settings=True)
    assert PermissionStore().get_policy().rules == []
    assert db.sensitive_integrity_check()["ok"]
    with db.connect() as conn:
        conn.execute("DELETE FROM permission_policies WHERE id = 'default'")
    with pytest.raises(db.SensitiveRecordIntegrityError):
        PermissionStore().get_policy()


def test_retention_removes_approval_presence_with_record():
    old = (datetime.now(UTC) - timedelta(days=60)).isoformat()
    task = Task(user_goal="old task", status=TaskPhase.COMPLETED, created_at=old)
    approval = Approval(task_id=task.id, message="old approval", created_at=old)
    db.upsert_model("tasks", task)
    db.upsert_model("approvals", approval)
    with db.connect() as conn:
        conn.execute("UPDATE tasks SET updated_at = ? WHERE id = ?", (old, task.id))
    cleanup_expired_task_details(retention_days=1, vacuum=False)
    assert db.fetch_one("approvals", approval.id) is None
    assert db.sensitive_integrity_check()["ok"]


def test_observe_scope_does_not_authorize_proactive_write(tmp_path):
    scope = store.save(AutonomyScope(label="observe", directories=[str(tmp_path)], allowed_actions=["observe"]))
    task = Task(user_goal="watch", metadata={"butler_matter_id": "matter", "butler_execution_ticket": "ticket"})
    store.save(
        ButlerMatter(
            id="matter",
            goal=task.user_goal,
            source="event",
            status=MatterStatus.RUNNING,
            active_task_id=task.id,
            execution_ticket="ticket",
            scope_versions={scope.id: scope.version},
        )
    )
    tool = SimpleNamespace(name="file.write_text", risk_level=RiskLevel.R2_REVERSIBLE_MODIFY)
    with pytest.raises(ButlerExecutionDeferred):
        validate_task_action(task, tool, {"path": str(tmp_path / "new.txt")}, live=True)


def test_public_matter_preserves_interaction_contract():
    parent = store.save(ButlerMatter(goal="parent"))
    matter = ButlerMatter(
        goal="ask",
        interaction=MatterInteraction(message="when?", fields=["date"]),
        constraints={
            "coordination_request": {"parents": [parent.id], "children": [], "versions": {parent.id: parent.version}}
        },
    )
    visible = public_matter(matter)
    assert visible["interaction"]["fields"] == ["date"]
    assert visible["constraints"]["coordination_request"]["parents"] == [parent.id]


def test_privacy_erase_clears_fts_content(tmp_path):
    from app.indexer.fts_index import FTSIndex

    document = tmp_path / "private.txt"
    document.write_text("ERASE_FTS_PRIVATE_MARKER", encoding="utf-8")
    FTSIndex().index_file(document, [str(tmp_path)])
    db.erase_local_user_data()
    with db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM document_chunks_fts").fetchone()[0] == 0
    assert b"ERASE_FTS_PRIVATE_MARKER" not in db.db_path().read_bytes()


def test_expired_approval_finishes_waiting_run_and_step():
    from app.core.schemas import ExecutionStage, Plan, PlanStep, Run, RunPhase, StepStatus
    from app.services import mobile_pairing_service, run_service

    task = Task(user_goal="approval", status=TaskPhase.EXECUTION, execution_stage=ExecutionStage.AWAITING_APPROVAL)
    step = PlanStep(
        task_id=task.id,
        agent_name="FileAgent",
        tool_name="file.write_text",
        description="write",
        status=StepStatus.WAITING_USER_APPROVAL,
    )
    plan = Plan(task_id=task.id, goal=task.user_goal, steps=[step])
    run = Run(message=task.user_goal, task_id=task.id, phase=RunPhase.AWAITING_APPROVAL)
    approval = Approval(
        task_id=task.id,
        step_id=step.id,
        message="expired",
        expires_at=(datetime.now(UTC) - timedelta(minutes=1)).isoformat(),
    )
    for table, model in (("tasks", task), ("plans", plan), ("runs", run), ("approvals", approval)):
        db.upsert_model(table, model)
    mobile_pairing_service._expire_stale_approval_records()
    assert db.fetch_one("plans", plan.id)["steps"][0]["status"] == "denied"
    assert run_service.resume_run(run.id).phase == RunPhase.DENIED
    assert db.fetch_one("tasks", task.id)["status"] == "denied"


def test_reviewed_failure_does_not_count_as_success_or_block_replanning():
    from app.core.schemas import ToolCall, ToolResult
    from app.services.butler_evidence import reviewed_results, uncertain_execution

    task = Task(user_goal="failed write")
    db.upsert_model("tasks", task)
    matter = ButlerMatter(goal=task.user_goal, task_ids=[task.id])
    call = ToolCall(
        task_id=task.id,
        step_id="step",
        tool_name="file.write_text",
        status="committed",
        risk_level=RiskLevel.R2_REVERSIBLE_MODIFY,
        dry_run=False,
    )
    result = ToolResult(
        tool_call_id=call.id,
        ok=False,
        output={"error": "Path is not a file."},
        runtime_review_completed=True,
        runtime_review_id="review",
        runtime_review_verdict="allow",
    )
    db.upsert_model("tool_calls", call)
    db.upsert_model("tool_results", result)
    assert not reviewed_results(matter)
    assert not uncertain_execution(matter)
    result.output["outcome_unknown"] = True
    db.upsert_model("tool_results", result)
    assert uncertain_execution(matter)


@pytest.mark.parametrize("kind", ["Submit", "submit ", "submit-form"])
def test_submission_aliases_share_business_identity(kind):
    from app.services.butler_effects import submission_identity

    canonical = {"url": "https://example.com/form", "selector": "#form", "fields": {"ref": "123"}}
    alias = {**canonical, "url": "https://EXAMPLE.com:443/a/../form#receipt", "kind": kind}
    assert submission_identity("browser.act", {"action": alias}) == submission_identity(
        "browser.submit_form", canonical
    )


def test_old_signed_submission_keeps_duplicate_protection_after_upgrade():
    from app.core.schemas import ToolCall
    from app.policy.approval_binding import hmac_digest
    from app.services.butler_effects import repeated_submission, submission_binding_hmac

    task = Task(user_goal="old submission")
    db.upsert_model("tasks", task)
    matter = ButlerMatter(goal=task.user_goal, task_ids=[task.id])
    original = {"url": "https://EXAMPLE.com:443/a/../form", "selector": "#form", "fields": {"ref": "123"}}
    call = ToolCall(
        task_id=task.id,
        step_id="step",
        tool_name="browser.submit_form",
        status="committed",
        dry_run=False,
        risk_level=RiskLevel.R3_DESTRUCTIVE_OR_SYSTEM,
        args={**original, "selector": "***", "fields": {"ref": "***"}},
        submission_digest=hmac_digest(original, prefix="browser-submission/v1"),
    )
    call.submission_binding_hmac = submission_binding_hmac(call)
    db.upsert_model("tool_calls", call)
    assert repeated_submission(matter, call.tool_name, {**original, "url": "https://example.com/form"})


def test_docx_redaction_covers_headers_footers_and_textboxes(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement

    from app.services.document_intelligence_service import _apply_redaction_docx

    path = tmp_path / "private.docx"
    document = Document()
    document.add_paragraph("alice@example.com")
    document.sections[0].header.paragraphs[0].text = "alice@example.com"
    document.sections[0].footer.paragraphs[0].text = "alice@example.com"
    textbox = OxmlElement("w:txbxContent")
    paragraph = OxmlElement("w:p")
    for text in ("alice@", "example.com"):
        run = OxmlElement("w:r")
        node = OxmlElement("w:t")
        node.text = text
        run.append(node)
        paragraph.append(run)
    textbox.append(paragraph)
    document.paragraphs[0]._p.append(textbox)
    document.save(path)
    _apply_redaction_docx(path, None)
    from zipfile import ZipFile

    with ZipFile(path) as package:
        assert not any(b"alice@" in package.read(name) for name in package.namelist() if name.endswith(".xml"))


def test_macro_workbook_edit_preserves_vba_part(tmp_path):
    from zipfile import ZipFile

    from openpyxl import Workbook

    from app.services.document_intelligence_service import edit_xlsx

    source, target = tmp_path / "source.xlsx", tmp_path / "macro.xlsm"
    Workbook().save(source)
    with ZipFile(source) as original, ZipFile(target, "w") as package:
        for name in original.namelist():
            package.writestr(name, original.read(name))
        package.writestr("xl/vbaProject.bin", b"SYNTHETIC_VBA_PART")
    assert edit_xlsx(target, sheet="Sheet", cell="A1", value="updated", dry_run=False)["ok"]
    with ZipFile(target) as package:
        assert package.read("xl/vbaProject.bin") == b"SYNTHETIC_VBA_PART"


@pytest.mark.parametrize("selector", ["#ssn", "#pin", "#totp", '[autocomplete="cc-number"]'])
def test_browser_read_boundary_rejects_sensitive_fields(selector):
    from app.services.browser_activity_runtime import BrowserActivityRuntime
    from app.tools.browser_tools import read_page

    assert not read_page({"url": "https://example.test/", "selector": selector}, {})["ok"]
    review = BrowserActivityRuntime._review_action(None, {"kind": "observe", "selectors": {"secret": selector}}, {})
    assert not review["ok"]
    assert str(review["verdict"]) == "deny"


@pytest.mark.asyncio
async def test_guardian_schedule_preserves_timezone():
    from app.services.guardian_scheduler import GuardianScheduler
    from app.services.scheduler_service import Scheduler, _next_run

    schedule = Scheduler().schedule("0 9 * * *", "morning", timezone="Asia/Shanghai")
    due = datetime(2026, 10, 19, 1, 0, tzinfo=UTC)
    schedule.next_run_at = due.isoformat()
    db.upsert_model("scheduled_tasks", schedule)
    guardian = GuardianScheduler(full_backend_available=AsyncMock(return_value=False))
    assert await guardian.tick(now=due) == [schedule.id]
    assert db.fetch_one("scheduled_tasks", schedule.id)["next_run_at"] == _next_run(
        schedule.cron, base=due, timezone="Asia/Shanghai"
    )


def test_desktop_screenshot_artifact_survives_bridge_and_adapter(monkeypatch):
    from app.services.browser_host_bridge_service import _sanitize_result
    from app.services.desktop_browser_adapter import DesktopBrowserActivityAdapter

    url = "file:///C:/Temp/lengrvis-browser-screenshots-0123456789abcdef/00000000-0000-0000-0000-000000000001.png"
    result = _sanitize_result({"ok": True, "event": {"screenshot_url": url}})
    adapter = DesktopBrowserActivityAdapter()
    monkeypatch.setattr(adapter, "_ensure_open", lambda *args: None)
    monkeypatch.setattr(adapter, "_request", lambda *args: result)
    assert adapter.perform(SimpleNamespace(current_url="https://example.test/"), {}, {})["screenshot_url"] == url
    assert (
        _sanitize_result({"ok": True, "event": {"screenshot_url": "file:///C:/private.png"}})["event"]["screenshot_url"]
        != "file:///C:/private.png"
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("action", ["pause", "cancel"])
async def test_terminal_matter_control_is_idempotent(action):
    from app.services.butler_service import ButlerCoordinator

    matter = store.save(ButlerMatter(goal="finished", status=MatterStatus.COMPLETED))
    assert (await ButlerCoordinator().control_matter(matter.id, action)).status == MatterStatus.COMPLETED
    assert store.get(ButlerMatter, matter.id).version == matter.version


def test_ui_action_route_removes_only_dispatch_keys(monkeypatch):
    from app.api import routes_ui_automation as route
    from app.core.schemas import SafetyVerdict
    from app.tools.ui_automation_contracts import validate_arguments

    tool = route._tool_definition("ui_automation.click")
    calls = []

    def execute(args, context):
        validate_arguments("click", args, context)
        calls.append(args)
        return {"ok": True, "dry_run": True}

    monkeypatch.setattr(tool, "execute", execute)
    monkeypatch.setattr(
        route,
        "_review_tool_call",
        lambda *args: SimpleNamespace(verdict=SafetyVerdict.ALLOW, risk_level=tool.risk_level),
    )
    route.action({"kind": "click", "automation_id": "target", "dry_run": True})
    assert len(calls) == 1
    assert calls[0] == {"automation_id": "target", "dry_run": True}


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", ["exception", "max_turns"])
async def test_engine_failure_finishes_linked_task(failure):
    from app.core.schemas import Run, RunPhase
    from app.orchestration.execution_models import RunPhase as EnginePhase
    from app.orchestration.execution_models import RunState
    from app.services import run_service

    task = Task(user_goal="failed engine", status=TaskPhase.PLANNING)
    run = Run(message=task.user_goal, task_id=task.id, phase=RunPhase.RUNNING)
    db.upsert_model("tasks", task)
    db.upsert_model("runs", run)
    state = RunState(
        run_id=run.id,
        task_id=task.id,
        engine="os",
        phase=EnginePhase.RUNNING,
        turn_count=1 if failure == "max_turns" else 0,
    )
    router = SimpleNamespace(max_turns=1, run_turn=AsyncMock(side_effect=RuntimeError("engine failed")))
    await run_service._run_engine_loop_body(run.id, router, state, stop_event=None, bridge_task=None)
    assert run_service.get_run(run.id).phase == RunPhase.FAILED
    assert db.fetch_one("tasks", task.id)["status"] == "failed"


def test_fts_filters_authorization_before_limiting_results(tmp_path):
    from app.indexer.fts_index import FTSIndex

    allowed, blocked = tmp_path / "allowed", tmp_path / "blocked"
    allowed.mkdir()
    blocked.mkdir()
    with db.connect() as conn:
        for file_id, directory in (("allowed", allowed), ("blocked", blocked)):
            path = str(directory / "notes.txt")
            conn.execute(
                "INSERT INTO indexed_files(id,normalized_path,data,sha256,name,extension,size,modified_at,indexed_at) "
                "VALUES(?,?,?,?,?,?,?,?,?)",
                (file_id, path, json.dumps({"path": path}), file_id, "notes.txt", ".txt", 1, "", ""),
            )
        for index in range(1002):
            file_id = "allowed" if index == 1001 else "blocked"
            path = str((allowed if file_id == "allowed" else blocked) / "notes.txt")
            conn.execute(
                "INSERT INTO document_chunks(id,file_id,chunk_index,text,data) VALUES(?,?,?,?,?)",
                (str(index), file_id, index, "needle xy", "{}"),
            )
            conn.execute(
                "INSERT INTO document_chunks_fts(file_id,path,text,chunk_id) VALUES(?,?,?,?)",
                (file_id, path, "needle xy", str(index)),
            )
    for query in ("needle", "xy"):
        results = FTSIndex().search(query, limit=1, allowed_directories=[str(allowed)])
        assert len(results) == 1
        assert results[0]["file_id"] == "allowed"
