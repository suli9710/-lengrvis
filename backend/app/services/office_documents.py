"""Create and independently read back bounded macro-free Office files."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile


def _text(value) -> str:
    return "" if value is None else str(value)


def render_workbook(content: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "整理数据"
    rows = [content["headers"] + ["来源网址", "原文依据"]]
    rows += [row["values"] + [row["source_url"], row["quote"]] for row in content["rows"]]
    for row_index, values in enumerate(rows, 1):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            # A web string beginning with '=' is data, never an executable formula.
            if isinstance(value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sources = workbook.create_sheet("资料来源")
    sources.append(["来源网址", "网页标题", "采集时间"])
    for source in content["sources"]:
        sources.append([source["url"], source["title"], source["observed_at"]])
    for worksheet in workbook:
        for row in worksheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243746")
        for column in worksheet.columns:
            worksheet.column_dimensions[column[0].column_letter].width = 28
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def report_projection(content: dict) -> dict:
    paragraphs = [content["title"], "摘要", content["summary"]]
    for section in content["sections"]:
        paragraphs.extend([section["heading"], section["text"]])
    paragraphs.extend(["整理数据", "资料来源"])
    paragraphs.extend(
        f"[{index}] {source['title']}\n{source['url']}\n采集时间：{source['observed_at']}"
        for index, source in enumerate(content["sources"], 1)
    )
    table = [content["headers"] + ["来源网址"]]
    table += [[_text(v) for v in row["values"]] + [row["source_url"]] for row in content["rows"]]
    return {"paragraphs": paragraphs, "tables": [table]}


def render_report(content: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.styles["Normal"].font.name = "Microsoft YaHei"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(content["title"], 0)
    document.add_heading("摘要", 1)
    document.add_paragraph(content["summary"])
    for section in content["sections"]:
        document.add_heading(section["heading"], 1)
        document.add_paragraph(section["text"])
    document.add_heading("整理数据", 1)
    values = report_projection(content)["tables"][0]
    table = document.add_table(rows=1, cols=len(values[0]))
    table.style = "Light Shading Accent 1"
    for index, row in enumerate(values):
        cells = table.rows[0].cells if index == 0 else table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value
    document.add_heading("资料来源", 1)
    for paragraph in report_projection(content)["paragraphs"][-len(content["sources"]) :]:
        document.add_paragraph(paragraph)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def inspect_document(data: bytes, kind: str) -> dict:
    # These are owned artifacts, but bound parsing also protects changed files.
    if len(data) > 8 * 1024 * 1024:
        raise ValueError("Office artifact exceeds the file size limit.")
    with ZipFile(BytesIO(data)) as package:
        if len(package.infolist()) > 2000 or sum(i.file_size for i in package.infolist()) > 32 * 1024 * 1024:
            raise ValueError("Office artifact exceeds the expanded size limit.")
    if kind == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if workbook.sheetnames != ["整理数据", "资料来源"]:
                raise ValueError("Workbook sheets changed.")
            sheets = {}
            for worksheet in workbook:
                if worksheet.max_row > 501 or worksheet.max_column > 32:
                    raise ValueError("Workbook dimensions exceed the bounded result.")
                rows = []
                for row in worksheet:
                    if any(cell.data_type == "f" for cell in row):
                        raise ValueError("Unexpected formula in source data.")
                    rows.append([cell.value for cell in row])
                sheets[worksheet.title] = rows
            return sheets
        finally:
            workbook.close()
    from docx import Document

    document = Document(BytesIO(data))
    return {
        "paragraphs": [p.text for p in document.paragraphs],
        "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables],
    }


def expected_projection(content: dict, kind: str) -> dict:
    if kind == "docx":
        return report_projection(content)
    # Excel represents an empty literal string as an empty cell when reopened.
    normalize = lambda rows: [[None if value == "" else value for value in row] for row in rows]  # noqa: E731
    return {
        "整理数据": normalize(
            [content["headers"] + ["来源网址", "原文依据"]]
            + [r["values"] + [r["source_url"], r["quote"]] for r in content["rows"]]
        ),
        "资料来源": [["来源网址", "网页标题", "采集时间"]]
        + [[s["url"], s["title"], s["observed_at"]] for s in content["sources"]],
    }


def verify_document(path: Path, kind: str, content: dict) -> tuple[str, dict]:
    import hashlib

    if path.stat().st_size > 8 * 1024 * 1024:
        raise ValueError("Office artifact exceeds the file size limit.")
    data = path.read_bytes()
    projection = inspect_document(data, kind)
    if projection != expected_projection(content, kind):
        raise ValueError("成果内容与已保存的表格、章节或来源不一致，请核对。")
    return hashlib.sha256(data).hexdigest(), projection
