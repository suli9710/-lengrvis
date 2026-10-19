"""Reproduce DOCX redaction leaving header/footer PII untouched while reporting success."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

os.environ.setdefault("LENGRVIS_AUDIT_HMAC_SECRET", "review-repro-secret")
os.environ["LENGRVIS_DATA_DIR"] = tempfile.mkdtemp(prefix="lengrvis-redact-")

from docx import Document  # noqa: E402

from app.services.document_intelligence_service import _apply_redaction_docx  # noqa: E402

path = Path(tempfile.mkdtemp(prefix="lengrvis-docx-")) / "memo.docx"
document = Document()
document.add_paragraph("Body contact: alice@example.com")
section = document.sections[0]
section.header.paragraphs[0].text = "Header contact: alice@example.com"
section.footer.paragraphs[0].text = "Footer contact: alice@example.com"
document.save(path)

_apply_redaction_docx(path, None)
document = Document(str(path))
print("body=", document.paragraphs[0].text)
print("header=", document.sections[0].header.paragraphs[0].text)
print("footer=", document.sections[0].footer.paragraphs[0].text)
if "alice@example.com" in document.sections[0].header.paragraphs[0].text:
    print("HEADER_PII_RETAINED")
